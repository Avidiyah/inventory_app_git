"""The Owner-only SMR audit workbook (spec 2026-10-23 D6): one summary row per
order with BEFORE / AFTER / DELTA triples, plus line-level sheets whose rows
are tagged by phase and share one header, archived orders included."""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from decimal import Decimal

import openpyxl

from app.domain import roles
from app.models import Item, User, WorkOrderItem
from app.services import auth
from app.services import work_orders as wos
from app.services import work_orders_smr as smr_service
from app.services import work_orders_smr_xlsx as xlsx


def _user(db, role=roles.ROLE_OWNER):
    user = User(username=f"u-{uuid.uuid4().hex[:8]}", first_name="Ola", last_name="Owner",
                password_hash=auth.hash_password("x"), role=role)
    db.add(user)
    db.flush()
    return user


def _sheet(numbers):
    head = "WORK ORDER,CREATOR,SERVICE TYPE,SCHEDULE DATE,COMPLETION DATE,ASSIGNED TO,FILED STATUS,LOCATION\r\n"
    rows = "".join(f"{n},Uri,SMR27 - Belfor,05/15/2026,05/21/2026,a,Closed,\"B:1,U 1,x\"\r\n" for n in numbers)
    return (head + rows).encode()


def _line(db, order, user, qty="2", price="1.50"):
    item = Item(barcode=f"BC-{uuid.uuid4().hex[:8]}", name="Rags", quantity=Decimal(100), location="Bay", price=Decimal(price))
    db.add(item)
    db.flush()
    line = WorkOrderItem(work_order_id=order.id, item_id=item.id, quantity=Decimal(qty), mode="retroactive", created_by_id=user.id)
    db.add(line)
    db.flush()
    db.refresh(order)
    return line


def _book(db):
    return openpyxl.load_workbook(io.BytesIO(xlsx.smr_workbook(db)))


def _rows(sheet):
    values = list(sheet.iter_rows(values_only=True))
    header = list(values[0])
    return header, [dict(zip(header, row)) for row in values[1:]]


def test_smr_export_sheets_and_phase_parity(db):
    owner = _user(db)
    number = f"SMR-{uuid.uuid4().hex[:8]}"
    order = wos.get_or_create_work_order(db, number=number, created_by_id=owner.id)
    _line(db, order, owner)
    smr_service.import_smr(db, csv_bytes=_sheet([number]), user=owner)
    order = wos.find_by_number(db, number)
    wos.add_work_order_labor(db, order.id, user=owner, technician_id=None, minutes=30)

    book = _book(db)
    assert book.sheetnames == ["Orders", "Material Lines", "Labor Entries", "Labor Sessions"]

    header, orders = _rows(book["Orders"])
    assert header[:8] == ["WORK ORDER", "CREATOR", "SERVICE TYPE", "SCHEDULE DATE", "COMPLETION DATE", "ASSIGNED TO", "FILED STATUS", "LOCATION"]
    for name in ("MATERIALS TOTAL", "SKILLED MINUTES", "LABOR TOTAL", "TOTAL"):
        assert {f"{name} BEFORE", f"{name} AFTER", f"{name} DELTA"} <= set(header)
    assert header[-6:] == ["EDITED", "SMR IMPORTED AT", "BASELINE AT", "LAST EDITED AT", "LAST EDITED BY", "ARCHIVED AT"]
    row = next(r for r in orders if r["WORK ORDER"] == number)
    assert row["CREATOR"] == "Uri"
    assert row["MATERIALS TOTAL BEFORE"] == 3.0 and row["MATERIALS TOTAL DELTA"] == 0
    assert row["SKILLED MINUTES DELTA"] == 30
    assert row["EDITED"] == "YES"
    assert row["LAST EDITED BY"] == "Ola Owner"
    assert row["STATUS"] == order.status

    header, lines = _rows(book["Material Lines"])
    mine = [r for r in lines if r["WORK ORDER"] == number]
    assert [r["PHASE"] for r in mine] == ["BEFORE", "AFTER"]
    before, after = mine
    assert {k: v for k, v in before.items() if k != "PHASE"} == {k: v for k, v in after.items() if k != "PHASE"}
    assert before["LINE TOTAL"] == 3.0

    header, labor = _rows(book["Labor Entries"])
    mine = [r for r in labor if r["WORK ORDER"] == number]
    assert [(r["PHASE"], r["TECHNICIAN"], r["MINUTES"]) for r in mine] == [("AFTER", "Skilled Technician", 30)]


def test_smr_export_removed_line_shows_only_before(db):
    owner = _user(db)
    number = f"SMR-{uuid.uuid4().hex[:8]}"
    order = wos.get_or_create_work_order(db, number=number, created_by_id=owner.id)
    line = _line(db, order, owner)
    smr_service.import_smr(db, csv_bytes=_sheet([number]), user=owner)
    db.delete(line)
    db.flush()

    book = _book(db)
    _, lines = _rows(book["Material Lines"])
    assert [r["PHASE"] for r in lines if r["WORK ORDER"] == number] == ["BEFORE"]
    _, orders = _rows(book["Orders"])
    row = next(r for r in orders if r["WORK ORDER"] == number)
    assert row["MATERIALS TOTAL DELTA"] == -3.0 and row["EDITED"] == "YES"


def test_smr_export_includes_archived_orders(db):
    owner = _user(db)
    number = f"SMR-{uuid.uuid4().hex[:8]}"
    smr_service.import_smr(db, csv_bytes=_sheet([number]), user=owner)
    wos.archive_work_order(db, wos.find_by_number(db, number).id, user=owner)

    _, orders = _rows(_book(db)["Orders"])
    row = next(r for r in orders if r["WORK ORDER"] == number)
    assert row["ARCHIVED AT"] is not None and row["EDITED"] == "NO"


# --- target match (spec 2026-10-23-smr-target-match) ----------------------

def _breakdown(rows):
    head = "WO#,Unit#,Closed Date,Labor Type,BELFOR HRS,Add HRS,Duplicate HRS,Hours,Labor $,OT Hours,OT $,Total Material Cost,Total L&M,\r\n"
    body = "".join(f'{n},Young Hall 108,05/21/26,SL,,,,{h},$0.00,,$0.00,"{m}",$0.00,\r\n' for n, h, m in rows)
    return (head + body).encode()


def test_smr_export_target_and_gap_columns(db):
    owner = _user(db)
    number, bare = [f"SMR-{uuid.uuid4().hex[:8]}" for _ in range(2)]
    smr_service.import_smr(db, csv_bytes=_sheet([number, bare]), user=owner)
    smr_service.import_smr_targets(db, csv_bytes=_breakdown([(number, "2", "$67.16")]), user=owner)
    order = wos.find_by_number(db, number)
    _line(db, order, owner, qty="2", price="20.00")  # 40.00 base, 46.00 marked
    wos.add_work_order_labor(db, order.id, user=owner, technician_id=None, minutes=60)

    header, rows = _rows(_book(db)["Orders"])

    assert header[-6:] == ["EDITED", "SMR IMPORTED AT", "BASELINE AT", "LAST EDITED AT", "LAST EDITED BY", "ARCHIVED AT"]
    assert header.index("TARGET HOURS") > header.index("TOTAL DELTA")
    row = next(r for r in rows if r["WORK ORDER"] == number)
    assert (row["TARGET HOURS"], row["TARGET MATERIAL $"], row["MATERIALS +15% AFTER"]) == (2.0, 67.16, 46.0)
    assert (row["GAP HOURS"], row["GAP MATERIAL $"], row["MATCHED"]) == (-1.0, -21.16, "NO")
    assert row["BREAKDOWN Unit#"] == "Young Hall 108" and row["TARGET IMPORTED AT"] is not None
    empty = next(r for r in rows if r["WORK ORDER"] == bare)
    assert empty["TARGET HOURS"] is None and empty["MATCHED"] is None
