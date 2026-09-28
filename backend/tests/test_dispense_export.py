"""Weekly export: dispenses since the previous export, over real HTTP."""

import io
import os
import sys
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import get_db
from app.main import app
from app.models import DispenseExport, Item, Transaction, User
from app.services import auth as auth_service
from app.services import dispense_export_xlsx

T0 = datetime(2026, 10, 5, 12, tzinfo=timezone.utc)


def _fresh(db):
    """Real local exports would move the window; the test transaction's
    rollback restores them."""
    db.query(DispenseExport).delete()
    db.commit()


def _user(db, role="techfm_oa"):
    user = User(username=f"u-{uuid.uuid4().hex[:10]}",
                password_hash=auth_service.hash_password("hunter2"), role=role)
    db.add(user)
    db.commit()
    return user


def _request(db, user, method, path):
    token = auth_service.create_session(db, user)
    app.dependency_overrides[get_db] = lambda: db
    try:
        with TestClient(app) as client:
            client.cookies.set("session", token)
            return client.request(method, path)
    finally:
        del app.dependency_overrides[get_db]


def _item_with(db, *moves):
    item = Item(barcode=f"DX-{uuid.uuid4().hex[:10]}", name=f"Dispensed {uuid.uuid4().hex[:6]}",
                quantity=Decimal("10"), location="Bay 1")
    db.add(item)
    db.flush()
    for kind, qty, at in moves:
        db.add(Transaction(item_id=item.id, transaction_type=kind,
                           quantity=Decimal(qty), created_at=at))
    db.commit()
    return item


def _sheet_rows(content):
    sheet = load_workbook(io.BytesIO(content)).active
    return {row[0]: row[6] for row in sheet.iter_rows(min_row=5, values_only=True) if row[0]}


def test_first_export_starts_at_midnight_central_on_launch_day(db):
    _fresh(db)
    dispense_export_xlsx.create_export(db, user=_user(db), now=T0)
    run = db.query(DispenseExport).one()
    assert run.window_start == datetime(2026, 9, 28, 5, tzinfo=timezone.utc)  # 00:00 CDT


def test_each_export_starts_where_the_last_ended_and_counts_only_dispenses(db):
    _fresh(db)
    user = _user(db)
    dispense_export_xlsx.create_export(db, user=user, now=T0)
    item = _item_with(
        db,
        ("dispense", "2", T0 + timedelta(hours=1)),
        ("dispense", "3", T0 + timedelta(hours=2)),
        ("stock", "50", T0 + timedelta(hours=3)),
        ("dispense", "7", T0 - timedelta(hours=1)),  # already printed last time
    )
    quiet = _item_with(db, ("adjust", "-1", T0 + timedelta(hours=1)))

    content, _ = dispense_export_xlsx.create_export(db, user=user, now=T0 + timedelta(days=7))

    rows = _sheet_rows(content)
    assert rows[item.name] == 5
    assert quiet.name not in rows


def test_reprint_returns_the_last_window_without_advancing(db):
    _fresh(db)
    user = _user(db)
    dispense_export_xlsx.create_export(db, user=user, now=T0)
    item = _item_with(db, ("dispense", "4", T0 + timedelta(hours=1)))
    dispense_export_xlsx.create_export(db, user=user, now=T0 + timedelta(days=7))
    count = db.query(DispenseExport).count()

    response = _request(db, user, "GET", "/items/dispense-exports/latest")

    assert response.status_code == 200
    assert _sheet_rows(response.content)[item.name] == 4
    assert db.query(DispenseExport).count() == count


def test_routes_are_gated_at_techfm_oa_and_reprint_404s_before_any_export(db):
    _fresh(db)
    supervisor = _user(db, "supervisor")
    assert _request(db, supervisor, "POST", "/items/dispense-exports").status_code == 403
    assert _request(db, supervisor, "GET", "/items/dispense-exports/latest").status_code == 403

    oa = _user(db)
    assert _request(db, oa, "GET", "/items/dispense-exports/latest").status_code == 404
    response = _request(db, oa, "POST", "/items/dispense-exports")
    assert response.status_code == 200
    assert "weekly-export_" in response.headers["content-disposition"]
    assert db.query(DispenseExport).count() == 1
