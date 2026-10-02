"""`GET /items/export`: the print workbook, over real HTTP."""

import io
import os
import sys
import uuid
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import get_db
from app.main import app
from app.models import Item, User
from app.services import auth as auth_service
from app.services import user_requests as request_service


def _export(db, role, path="/items/export"):
    user = User(
        username=f"u-{uuid.uuid4().hex[:10]}",
        password_hash=auth_service.hash_password("hunter2"),
        role=role,
    )
    db.add(user)
    db.commit()
    token = auth_service.create_session(db, user)
    app.dependency_overrides[get_db] = lambda: db
    try:
        with TestClient(app) as client:
            client.cookies.set("session", token)
            return client.get(path)
    finally:
        del app.dependency_overrides[get_db]


def test_export_is_gated_at_techfm_oa(db):
    assert _export(db, "supervisor").status_code == 403


def test_export_holds_items_and_every_request_tab(db):
    name = f"Export widget {uuid.uuid4().hex[:6]}"
    db.add(Item(barcode=f"BC-{uuid.uuid4().hex[:10]}", name=name,
                quantity=Decimal("3"), low_stock_threshold=Decimal("5"), location="Bay 1", notes={"color": "red"}))
    for text in ("Mystery pipe", "Settled pipe"):
        request_service.create_catalogue_request(
            db, searched_text=text, quantity=Decimal("2"), note=None,
            work_order_id=None, work_order_number=None, source="find_item", created_by_id=None,
        )
    db.flush()
    settled = next(r for r in request_service.list_user_requests(db) if
                   (r.details or {}).get("searched_text") == "Settled pipe")
    settled.status = request_service.STATUS_RESOLVED
    db.commit()

    response = _export(db, "techfm_oa")

    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == [
        "Items", "Material requests", "Catalogue requests",
        "Stock recounts", "Missing price - link",
    ]
    item_names = [row[0] for row in book["Items"].iter_rows(min_row=5, values_only=True)]
    assert name in item_names
    requests = [row[2] for row in book["Catalogue requests"].iter_rows(min_row=5, values_only=True)]
    assert "Mystery pipe" in requests
    assert "Settled pipe" not in requests  # resolved: nothing left to check
    assert book["Items"].print_title_rows == "$4:$4"
    assert "low stock" in book["Items"]["A2"].value
    low_rows = [r for r in book["Items"].iter_rows(min_row=5) if r[3].value == "LOW"]
    assert name in [r[0].value for r in low_rows]
    assert all(r[0].font.color.rgb.endswith("C8102E") for r in low_rows)


def test_labels_page_draws_svg_barcodes_and_escapes_names(db):
    barcode = f"LBL-{uuid.uuid4().hex[:8].upper()}"
    db.add(Item(barcode=barcode, name="Pipe <1in> & cap", quantity=Decimal("1"), location="Bay 1"))
    db.add(Item(barcode=f"lower-{uuid.uuid4().hex[:8]}", name="Lowercase code",
                quantity=Decimal("1"), location="Bay 1"))
    db.commit()

    assert _export(db, "supervisor", "/items/labels").status_code == 403
    response = _export(db, "techfm_oa", "/items/labels")

    assert response.status_code == 200
    assert '<svg class="bars"' in response.text
    assert f'<div class="code">{barcode}</div>' in response.text
    assert "Pipe &lt;1in&gt; &amp; cap" in response.text
    assert '<div class="no-bars">' in response.text  # lowercase: no Code 39


def test_labels_page_moves_the_parenthetical_to_its_own_line(db):
    tag = uuid.uuid4().hex[:8].upper()
    db.add(Item(barcode=f"LBL-{tag}", name=f'Blinds {tag} (35" x 72")',
                quantity=Decimal("1"), location="Bay 1"))
    db.commit()

    page = _export(db, "techfm_oa", f"/items/labels?barcode=LBL-{tag}").text

    assert (f'<div class="name">Blinds {tag}</div>'
            '<div class="name-sub">(35&quot; x 72&quot;)</div>') in page


def test_labels_page_narrows_to_a_search_or_chosen_barcodes(db):
    tag = uuid.uuid4().hex[:8].upper()
    for n in ("A", "B"):
        db.add(Item(barcode=f"LBL-{tag}-{n}", name=f"Narrow {tag} {n}",
                    quantity=Decimal("1"), location="Bay 1"))
    db.commit()

    searched = _export(db, "techfm_oa", f"/items/labels?q=Narrow {tag}").text
    assert f">LBL-{tag}-A<" in searched and f">LBL-{tag}-B<" in searched
    chosen = _export(db, "techfm_oa", f"/items/labels?barcode=LBL-{tag}-B").text
    assert f">LBL-{tag}-A<" not in chosen and f">LBL-{tag}-B<" in chosen


def test_labels_page_narrows_to_a_location_and_lists_every_location(db):
    tag = uuid.uuid4().hex[:8].upper()
    db.add(Item(barcode=f"LBL-{tag}-IN", name="In", quantity=Decimal("1"), location=f"Shelf {tag}"))
    db.add(Item(barcode=f"LBL-{tag}-OUT", name="Out", quantity=Decimal("1"), location=f"Bay {tag}"))
    db.commit()

    page = _export(db, "techfm_oa", f"/items/labels?location=Shelf {tag}").text

    assert f">LBL-{tag}-IN<" in page and f">LBL-{tag}-OUT<" not in page
    assert f'<option value="Shelf {tag}" selected>' in page
    assert f'<option value="Bay {tag}">' in page


def test_labels_page_has_size_sliders_up_to_the_printable_sheet(db):
    page = _export(db, "techfm_oa", "/items/labels").text

    assert 'data-var="w" min="1" max="7.5" step="0.05" value="2.5"' in page
    assert 'data-var="h" min="1" max="10" step="0.05" value="3"' in page
    assert '<script type="module" src="/static/labels.js"></script>' in page


def test_code39_modules_round_trip_through_the_decoder():
    from PIL import Image, ImageDraw

    from app.services.barcode_labels import _code39_modules
    from app.services.barcodes import decode_image

    text = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ-. $/+%"  # every table entry
    modules = "0" * 10 + _code39_modules(text) + "0" * 10
    image = Image.new("L", (len(modules) * 3, 80), 255)
    draw = ImageDraw.Draw(image)
    for i, m in enumerate(modules):
        if m == "1":
            draw.rectangle([i * 3, 0, i * 3 + 2, 79], fill=0)
    png = io.BytesIO()
    image.save(png, format="PNG")

    assert [(m.text, m.format) for m in decode_image(png.getvalue())] == [(text, "CODE_39")]
    assert _code39_modules("lower") is None and _code39_modules("A*B") is None
