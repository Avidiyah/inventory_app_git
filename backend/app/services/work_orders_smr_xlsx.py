"""The SMR27 audit workbook (spec 2026-10-23 §Export).

Four sheets. `Orders` is one row per SMR order: the vendor sheet's raw
columns, the operational export's columns, a BEFORE / AFTER / DELTA triple
for every number analytics will sum, then the audit trail. The three line
sheets carry one row per material line, labor entry and labor session, each
tagged `PHASE` = BEFORE (from the frozen baseline) or AFTER (the live order
through the same serializer), so the two phases are column-identical.
Money and minutes are written as numbers, never strings."""

from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.domain import work_orders as wo
from app.domain import work_orders_smr as smr
from app.models import User, WorkOrder
from app.services import _xlsx_theme as theme
from app.services import work_orders as wos
from app.services import work_orders_smr as smr_service

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
SMR_XLSX_FILENAME = "smr27-audit-export.xlsx"
TAB_COLOR = "000000"

BEFORE, AFTER = "BEFORE", "AFTER"

# Every number in a snapshot's `totals`, plus the line count, as (label, key).
_NUMERIC = (
    ("MATERIAL LINES", None),
    ("MATERIALS TOTAL", "materials_total"),
    ("REAL MINUTES", "real_minutes"),
    ("SKILLED MINUTES", "skilled_minutes"),
    ("REAL BILLED MINUTES", "real_billed_minutes"),
    ("SKILLED BILLED MINUTES", "skilled_billed_minutes"),
    ("LABOR TOTAL", "labor_total"),
    ("TOTAL", "total"),
)
# The client's target and the gap to it (spec 2026-10-23-smr-target-match);
# all blank for an order without a Breakdown row.
_TARGET = (
    "MATERIALS +15% AFTER", "TARGET HOURS", "TARGET LABOR $", "TARGET MATERIAL $", "TARGET TOTAL",
    "GAP HOURS", "GAP MATERIAL $", "MATCHED", "TARGET IMPORTED AT",
)
_BREAKDOWN_PREFIX = "BREAKDOWN "
_TRAIL = ("EDITED", "SMR IMPORTED AT", "BASELINE AT", "LAST EDITED AT", "LAST EDITED BY", "ARCHIVED AT")

MATERIAL_HEADERS = ("WORK ORDER", "PHASE", "ITEM", "BARCODE", "QUANTITY", "BILLABLE QUANTITY", "UNIT PRICE", "LINE TOTAL", "CREATED AT")
LABOR_HEADERS = ("WORK ORDER", "PHASE", "TECHNICIAN", "MINUTES", "RECORDED BY", "CREATED AT")
SESSION_HEADERS = ("WORK ORDER", "PHASE", "TECHNICIAN", "STARTED AT", "ENDED AT", "MINUTES", "AUTO CLOSED", "REVIEWED AT")


def _num(value) -> Optional[float]:
    """Snapshot numbers are strings (Decimal-safe JSON); Excel wants numbers."""
    if value is None:
        return None
    return float(Decimal(value)) if isinstance(value, str) else float(value)


def _when(value) -> Optional[datetime]:
    if value is None:
        return None
    parsed = datetime.fromisoformat(value) if isinstance(value, str) else value
    # openpyxl refuses tz-aware datetimes; the ISO text already says UTC.
    return parsed.replace(tzinfo=None)


def _source_columns(orders: Iterable[WorkOrder]) -> list[str]:
    """The sheet's own headers across every baseline, in sheet order."""
    seen: set[str] = set()
    for order in orders:
        seen.update(((order.smr_baseline or {}).get("source_row") or {}).keys())
    return smr.sheet_column_order(seen)


def _target_columns(orders: Iterable[WorkOrder]) -> list[str]:
    """The Breakdown sheet's headers across every target, in its own order."""
    seen: set[str] = set()
    for order in orders:
        seen.update(((order.smr_target or {}).get("source_row") or {}).keys())
    known = [c for c in smr.TARGET_SHEET_COLUMNS if c in seen]
    return known + sorted(c for c in seen if c not in smr.TARGET_SHEET_COLUMNS)


def _target_values(order: WorkOrder, current: dict) -> list:
    """`_TARGET` values for one order; a row of None without a target."""
    target = smr_service.target_for(order, current)
    if target is None:
        return [None] * len(_TARGET)
    return [
        _num(target.current_materials_total),
        _num(target.hours),
        _num(target.labor_total),
        _num(target.materials_total),
        _num(target.total),
        _num(target.gap_hours),
        _num(target.gap_materials_total),
        "YES" if target.matched else "NO",
        _when((order.smr_target or {}).get("imported_at")),
    ]


# Operational-export columns the sheet already supplies by the same name.
_SHEET_OWNED = {"WORK ORDER", "LOCATION", "ASSIGNED TO", "SERVICE TYPE", "SCHEDULE DATE"}
_APP_HEADERS = [h for h in wo.EXPORT_HEADERS if h not in _SHEET_OWNED]


def _number_value(snap: dict, key: Optional[str]):
    return len(snap["material_lines"]) if key is None else snap["totals"][key]


def _last_edit(order: WorkOrder, baseline_at: Optional[datetime], names: dict) -> tuple[Optional[datetime], str]:
    """The latest labor entry or material line added after the baseline."""
    latest: Optional[datetime] = None
    who = ""
    for entry in order.labor_entries:
        if entry.created_at and (baseline_at is None or entry.created_at > baseline_at):
            if latest is None or entry.created_at > latest:
                latest, who = entry.created_at, (entry.recorded_by.full_name if entry.recorded_by else "")
    for line in order.items:
        if line.created_at and (baseline_at is None or line.created_at > baseline_at):
            if latest is None or line.created_at > latest:
                latest, who = line.created_at, names.get(line.created_by_id, "")
    return latest, who


def _orders_sheet(sheet: Worksheet, orders: list[WorkOrder], names: dict) -> None:
    source_columns = _source_columns(orders)
    target_columns = _target_columns(orders)
    headers = (
        source_columns
        + _APP_HEADERS
        + [f"{label} {phase}" for label, _ in _NUMERIC for phase in (BEFORE, AFTER, "DELTA")]
        + list(_TARGET)
        + [f"{_BREAKDOWN_PREFIX}{column}" for column in target_columns]
        + list(_TRAIL)
    )
    theme.setup_sheet(sheet, tab_color=TAB_COLOR, freeze="B2")
    theme.header_row(sheet, 1, headers)
    rows = []
    for order in orders:
        current = smr_service.current_snapshot(order)
        baseline = order.smr_baseline or current
        source = baseline.get("source_row") or {}
        baseline_at = _when(baseline.get("taken_at"))
        edited_at, edited_by = _last_edit(order, order.smr_imported_at, names)
        numbers = []
        for _, key in _NUMERIC:
            b, a = _num(_number_value(baseline, key)), _num(_number_value(current, key))
            numbers += [b, a, a - b]
        breakdown = (order.smr_target or {}).get("source_row") or {}
        rows.append(
            [source.get(column) for column in source_columns]
            + [value for header, value in zip(wo.EXPORT_HEADERS, wos.export_row(order)) if header not in _SHEET_OWNED]
            + numbers
            + _target_values(order, current)
            + [breakdown.get(column) for column in target_columns]
            + [
                "YES" if smr.is_edited(order) else "NO",
                _when(order.smr_imported_at),
                baseline_at,
                _when(edited_at),
                edited_by,
                _when(order.archived_at),
            ]
        )
    theme.write_rows(sheet, 2, rows)


def _phase_rows(orders: list[WorkOrder], key: str, project) -> list[list]:
    rows: list[list] = []
    for order in orders:
        current = smr_service.current_snapshot(order)
        baseline = order.smr_baseline or current
        for phase, snap in ((BEFORE, baseline), (AFTER, current)):
            for record in snap[key]:
                rows.append([order.number, phase] + project(record))
    return rows


def _lines_sheet(sheet: Worksheet, headers: tuple[str, ...], rows: list[list]) -> None:
    theme.setup_sheet(sheet, tab_color=TAB_COLOR, freeze="C2")
    theme.header_row(sheet, 1, headers)
    theme.write_rows(sheet, 2, rows)


def smr_workbook(db: Session) -> bytes:
    orders = smr_service.smr_orders(db)
    creator_ids = {line.created_by_id for order in orders for line in order.items if line.created_by_id}
    names = {
        user.id: user.full_name
        for user in db.query(User).filter(User.id.in_(creator_ids)).all()
    } if creator_ids else {}

    workbook = Workbook()
    orders_sheet = workbook.active
    orders_sheet.title = "Orders"
    _orders_sheet(orders_sheet, orders, names)

    _lines_sheet(
        workbook.create_sheet("Material Lines"), MATERIAL_HEADERS,
        _phase_rows(orders, "material_lines", lambda r: [
            r["item_name"], r["barcode"], _num(r["quantity"]), _num(r["billable_quantity"]),
            _num(r["unit_price"]), _num(r["line_total"]), _when(r["created_at"]),
        ]),
    )
    _lines_sheet(
        workbook.create_sheet("Labor Entries"), LABOR_HEADERS,
        _phase_rows(orders, "labor_entries", lambda r: [
            r["technician"], r["minutes"], r["recorded_by"], _when(r["created_at"]),
        ]),
    )
    _lines_sheet(
        workbook.create_sheet("Labor Sessions"), SESSION_HEADERS,
        _phase_rows(orders, "labor_sessions", lambda r: [
            r["technician"], _when(r["started_at"]), _when(r["ended_at"]), r["minutes"],
            "YES" if r["auto_closed"] else "NO", _when(r["reviewed_at"]),
        ]),
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
